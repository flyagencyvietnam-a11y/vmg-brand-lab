"""Trích VMG Brand Book (docx) thành seed/brand-book.json. Không gõ lại bằng tay: mọi chữ lấy từ file nguồn.

    npm run seed:brandbook

Cần uv (https://docs.astral.sh/uv/). Nguồn: docs/source/VMG_Brand_Book_2026_v0.2_final.docx (bản dự thảo, chờ Ban Giám Đốc duyệt).

Brand Book là nguồn cho các mục đã có của Brand Lab (không có mục riêng). Đầu ra:
  meta    thông tin trang bìa (để ghi nguồn và phiên bản)
  brands  hồ sơ từng thương hiệu: bảng định vị mục 3.x cùng vai trò (1.3), vai nhân cách (2.5), định vị một câu, màu thẻ (4.2)
"""

import json
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/source/VMG_Brand_Book_2026_v0.2_final.docx"
OUT = ROOT / "seed/brand-book.json"

# Tên thương hiệu trong tài liệu -> slug trong seed/brands.json.
SLUGS = {
    "VMG": "vmg",
    "VMG English": "vmg-english",
    "VMG IELTS": "vmg-ielts",
    "VMG TESOL": "vmg-tesol",
    "VMG Tiếng Trung": "vmg-chinese",
    "VMP by VMG": "vmp",
    "VMT": "vmt",
    "UpLearn by VMG": "uplearn",
    "EduNext": "edunext",
    "VMG e-Academy": "vmg-e-academy",
}

def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def cell_paragraphs(cell) -> list[str]:
    return [clean(p.text) for p in cell.paragraphs if clean(p.text)]


def row_cells(row) -> list[list[str]]:
    """Các ô của một dòng, bỏ ô lặp do gộp ô."""
    cells: list[list[str]] = []
    seen = set()
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        cells.append(cell_paragraphs(c))
    return cells


def is_header_row(row) -> bool:
    """Dòng tiêu đề: mọi ô có chữ đều in đậm hoặc có tô nền."""
    any_text = False
    for c in row.cells:
        paras = [p for p in c.paragraphs if clean(p.text)]
        if not paras:
            continue
        any_text = True
        shaded = c._tc.find(qn("w:tcPr")) is not None and c._tc.find(qn("w:tcPr")).find(qn("w:shd")) is not None
        runs = [r for p in paras for r in p.runs if clean(r.text)]
        if not shaded and not (runs and all(r.bold for r in runs)):
            return False
    return any_text


def table_block(table: Table) -> dict:
    rows = table.rows
    cells = [row_cells(r) for r in rows]
    # Ô chú: bảng một ô.
    if len(rows) == 1 and len(cells[0]) == 1:
        return {"type": "callout", "paras": cells[0][0]}
    header = None
    body = cells
    if len(rows) > 1 and is_header_row(rows[0]):
        header = [" ".join(c) for c in cells[0]]
        body = cells[1:]
    return {"type": "table", "header": header, "rows": body}


def read_body(doc) -> list[dict]:
    """Chuỗi phần tử theo thứ tự: tiêu đề (h1, h2, h3), đoạn, danh sách, bảng."""
    items: list[dict] = []
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            text = clean(p.text)
            if not text:
                continue
            style = p.style.name
            if style == "Heading 1":
                items.append({"kind": "h1", "text": text})
            elif style == "Heading 2":
                items.append({"kind": "h2", "text": text})
            elif style == "Heading 3":
                items.append({"kind": "h3", "text": text})
            elif style == "List Paragraph":
                if items and items[-1]["kind"] == "ul":
                    items[-1]["items"].append(text)
                else:
                    items.append({"kind": "ul", "items": [text]})
            else:
                items.append({"kind": "p", "text": text})
        elif el.tag == qn("w:tbl"):
            items.append({"kind": "table", "block": table_block(Table(el, doc))})
    return items


def to_blocks(items: list[dict]) -> list[dict]:
    blocks: list[dict] = []
    for it in items:
        k = it["kind"]
        if k == "p":
            blocks.append({"type": "p", "text": it["text"]})
        elif k == "ul":
            blocks.append({"type": "ul", "items": it["items"]})
        elif k == "h3":
            blocks.append({"type": "h3", "text": it["text"]})
        elif k == "table":
            blocks.append(it["block"])
    return blocks


def section_no(title: str) -> str | None:
    m = re.match(r"^(\d\.\d)\s", title)
    return m.group(1) if m else None


def main() -> None:
    doc = docx.Document(str(SOURCE))
    items = read_body(doc)

    # Trang bìa: các đoạn trước "Mục lục" và mục đầu tiên.
    cover = [it["text"] for it in items[:8] if it["kind"] == "p"]
    meta = {
        "organisation": cover[0],
        "title": cover[1],
        "subtitle": cover[2],
        "tagline": cover[3],
        "version_line": cover[4],
        "author": cover[5],
        "submitted_to": cover[6],
    }

    # Cắt theo tiêu đề cấp 2 (mục 1.1...) và cấp 1 (phần, "Về tài liệu này").
    sections: dict[str, dict] = {}
    current: dict | None = None
    in_intro = False
    part_of: dict[str, str] = {}
    part_title = ""
    for it in items:
        if it["kind"] == "h1":
            part_title = it["text"]
            in_intro = it["text"].lower().startswith("về tài liệu")
            current = None
            continue
        if it["kind"] == "h2":
            no = section_no(it["text"])
            current = {"no": no, "title": re.sub(r"^\d\.\d\s+", "", it["text"]), "items": []}
            if no:
                sections[no] = current
                part_of[no] = part_title
            in_intro = False
            continue
        if in_intro:
            continue
        if current is not None:
            current["items"].append(it)

    # Hồ sơ thương hiệu.
    def table_of(no: str, index: int = 0) -> dict:
        tables = [it["block"] for it in sections[no]["items"] if it["kind"] == "table"]
        return tables[index]

    def by_name(table: dict) -> dict[str, list[list[str]]]:
        return {row[0][0]: row[1:] for row in table["rows"] if row and row[0]}

    roles = by_name(table_of("1.3"))  # [vai trò và đối tượng], [cách viết]
    persona = by_name(table_of("2.5"))  # [vai], [giọng nói]
    # Mục 3: bảng định vị một câu nằm trước mục 3.1, trong phần 3 chưa có tiêu đề cấp 2.
    one_liners: dict[str, str] = {}
    in_part3 = False
    for it in items:
        if it["kind"] == "h1":
            in_part3 = it["text"].startswith("PHẦN 3")
        elif in_part3 and it["kind"] == "table" and it["block"]["type"] == "table" and it["block"]["header"] and it["block"]["header"][0] == "Thương hiệu":
            for row in it["block"]["rows"]:
                one_liners[row[0][0]] = " ".join(row[1])
            break
    colours: dict[str, str] = {}
    for it in sections["4.2"]["items"]:
        if it["kind"] == "table" and it["block"]["type"] == "table" and it["block"]["header"] == ["Thương hiệu", "Màu thẻ"]:
            for row in it["block"]["rows"]:
                for name in re.split(r",\s*", row[0][0]):
                    colours[name] = " ".join(row[1])

    brands: dict[str, dict] = {}
    for no, sec in sections.items():
        if not no.startswith("3."):
            continue
        name = sec["title"]
        tbl = next(it["block"] for it in sec["items"] if it["kind"] == "table")
        profile = [{"label": row[0][0], "value": row[1]} for row in tbl["rows"] if len(row) >= 2 and row[0]]
        brands[SLUGS[name]] = {"name": name, "profile": profile}

    for name, slug in SLUGS.items():
        b = brands.setdefault(slug, {"name": name, "profile": []})
        r = roles.get(name)
        if r:
            b["architecture_role"] = " ".join(r[0])
            b["spelling"] = " ".join(r[1]) if len(r) > 1 else None
        p = persona.get(name)
        if p:
            b["persona_role"] = " ".join(p[0])
            b["persona_voice"] = " ".join(p[1]) if len(p) > 1 else None
        if name in one_liners:
            b["one_liner"] = one_liners[name]
        if name in colours:
            b["card_colour"] = colours[name]

    out = {
        "_source": "docs/source/VMG_Brand_Book_2026_v0.2_final.docx (bản dự thảo v0.2, Marketing soạn, trình Ban Giám Đốc). Sinh bằng `npm run seed:brandbook`, không sửa tay.",
        "_note": "Nội dung ở trạng thái dự thảo: nạp ở draft, chưa phát hành cho nhân viên cho đến khi Ban Giám Đốc phê duyệt (spec 9.1).",
        "meta": meta,
        "brands": brands,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"seed/brand-book.json: {len(brands)} thương hiệu", file=sys.stderr)


if __name__ == "__main__":
    main()
